
from __future__ import annotations

import base64
import io
import logging
import os
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any
import html as html_lib
from ai_summarizer_router import router as ai_summarizer_router
import fitz
import pdfplumber
from docx import Document
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response
from openpyxl import Workbook
from pdf2docx import Converter
from pptx import Presentation
from pptx.util import Inches

# Keep translation_router.py in the same backend directory.
from translation_router import router as translation_router


# ---------------------------------------------------------------------------
# Application configuration
# ---------------------------------------------------------------------------

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("pdf-toolkit")

app = FastAPI(
    title="PDF Toolkit API",
    version="1.3.0",
    description=(
        "PDF conversions, Office document conversions, "
        "HTML rendering, and PDF translation."
    ),
)

FRONTEND_ORIGINS = [
    origin.strip().rstrip("/")
    for origin in os.getenv(
        "FRONTEND_ORIGINS",
        (
            "http://localhost:3000,"
            "http://127.0.0.1:3000,"
            "https://pdftoolkit-opwa-p5v2apmor-usamawebdevelopers-projects.vercel.app"
        ),
    ).split(",")
    if origin.strip()
]

app.add_middleware(
    CORSMiddleware,
    allow_origins=FRONTEND_ORIGINS,
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

MAX_UPLOAD_BYTES = int(os.getenv("MAX_UPLOAD_MB", "30")) * 1024 * 1024
MAX_PDF_PAGES = 200
MAX_HTML_IMAGE_BYTES = 25 * 1024 * 1024

SUPPORTED_CONVERSIONS = {
    "pdf-to-word",
    "pdf-to-powerpoint",
    "pdf-to-excel",
    "pdf-to-html",
    "word-to-pdf",
    "powerpoint-to-pdf",
    "excel-to-pdf",
    "html-to-pdf",
}


# ---------------------------------------------------------------------------
# General helpers
# ---------------------------------------------------------------------------

def find_libreoffice() -> str | None:
    """Find the LibreOffice executable."""
    configured_path = os.getenv("LIBREOFFICE_PATH")

    candidates = [
        configured_path,
        shutil.which("libreoffice"),
        shutil.which("soffice"),
    ]

    if os.name == "nt":
        candidates.extend([
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        ])

    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return str(candidate)

    return None


def find_html_pdf_browser() -> str | None:
    """
    Find Google Chrome or Microsoft Edge for HTML-to-PDF rendering.

    HTML_PDF_BROWSER_PATH can specify a custom browser executable.
    """
    configured_path = os.getenv("HTML_PDF_BROWSER_PATH")

    candidates = [
        configured_path,
        shutil.which("chrome"),
        shutil.which("google-chrome"),
        shutil.which("google-chrome-stable"),
        shutil.which("chromium"),
        shutil.which("chromium-browser"),
        shutil.which("msedge"),
        shutil.which("microsoft-edge"),
    ]

    if os.name == "nt":
        program_files = os.getenv(
            "PROGRAMFILES",
            r"C:\Program Files",
        )
        program_files_x86 = os.getenv(
            "PROGRAMFILES(X86)",
            r"C:\Program Files (x86)",
        )
        local_app_data = os.getenv("LOCALAPPDATA", "")

        candidates.extend([
            str(
                Path(program_files)
                / "Google/Chrome/Application/chrome.exe"
            ),
            str(
                Path(program_files_x86)
                / "Google/Chrome/Application/chrome.exe"
            ),
            str(
                Path(program_files)
                / "Microsoft/Edge/Application/msedge.exe"
            ),
            str(
                Path(program_files_x86)
                / "Microsoft/Edge/Application/msedge.exe"
            ),
            str(
                Path(local_app_data)
                / "Google/Chrome/Application/chrome.exe"
            ),
            str(
                Path(local_app_data)
                / "Microsoft/Edge/Application/msedge.exe"
            ),
        ])

    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return str(candidate)

    return None


def safe_stem(
    filename: str | None,
    default: str = "document",
) -> str:
    """Return a filesystem-safe filename without its extension."""
    name = Path(filename or default).stem

    cleaned = "".join(
        char
        for char in name
        if char.isalnum() or char in ("-", "_", " ")
    ).strip()

    return cleaned[:100] or default


def file_response(
    content: bytes,
    filename: str,
    media_type: str = "application/octet-stream",
) -> Response:
    """Return a file as a downloadable response."""
    safe_filename = (
        filename
        .replace('"', "")
        .replace("\r", "")
        .replace("\n", "")
    )

    return Response(
        content=content,
        media_type=media_type,
        headers={
            "Content-Disposition": (
                f'attachment; filename="{safe_filename}"'
            ),
            "Content-Length": str(len(content)),
            "X-Content-Type-Options": "nosniff",
        },
    )


async def read_upload(
    upload: UploadFile,
    max_bytes: int = MAX_UPLOAD_BYTES,
) -> bytes:
    """Read an upload and enforce the configured size limit."""
    content = await upload.read(max_bytes + 1)

    if not content:
        raise HTTPException(
            status_code=400,
            detail="The uploaded file is empty.",
        )

    if len(content) > max_bytes:
        raise HTTPException(
            status_code=413,
            detail=(
                "File is too large. Maximum size is "
                f"{max_bytes // (1024 * 1024)} MB."
            ),
        )

    return content


def require_pdf(content: bytes) -> None:
    """Validate that the uploaded file is a readable PDF."""
    if not content.startswith(b"%PDF-"):
        raise HTTPException(
            status_code=400,
            detail="Invalid PDF file. Please upload a valid PDF.",
        )

    try:
        with fitz.open(stream=content, filetype="pdf") as document:
            if document.needs_pass:
                raise HTTPException(
                    status_code=400,
                    detail=(
                        "This PDF is password-protected. "
                        "Unlock it before converting."
                    ),
                )

            if document.page_count == 0:
                raise HTTPException(
                    status_code=400,
                    detail="The PDF contains no pages.",
                )

            if document.page_count > 500:
                raise HTTPException(
                    status_code=400,
                    detail="This PDF exceeds the 500-page safety limit.",
                )

    except HTTPException:
        raise
    except Exception as exc:
        logger.warning("PDF validation failed: %s", exc)
        raise HTTPException(
            status_code=400,
            detail=(
                "The PDF could not be opened. "
                "It may be damaged or invalid."
            ),
        ) from exc


def validate_conversion_slug(tool_slug: str) -> str:
    """Validate the requested conversion tool."""
    if tool_slug not in SUPPORTED_CONVERSIONS:
        raise HTTPException(
            status_code=404,
            detail=f"Unsupported conversion tool: {tool_slug}",
        )

    return tool_slug


def rgb_to_hex(color: int) -> str:
    """Convert a PyMuPDF integer RGB color to a CSS hex color."""
    color = max(0, min(int(color), 0xFFFFFF))
    return f"#{color:06x}"


def css_font_family(font_name: str) -> str:
    """Choose a sensible CSS font fallback from a PDF font name."""
    normalized = font_name.lower()

    if "mono" in normalized or "courier" in normalized:
        return "monospace"

    if "serif" in normalized or "times" in normalized:
        return "serif"

    return "Arial, Helvetica, sans-serif"


# ---------------------------------------------------------------------------
# PDF to Word
# ---------------------------------------------------------------------------

def convert_pdf_to_docx(pdf_bytes: bytes) -> bytes:
    """Convert a PDF document to Word."""
    with tempfile.TemporaryDirectory() as directory:
        source = Path(directory) / "input.pdf"
        target = Path(directory) / "output.docx"

        source.write_bytes(pdf_bytes)

        converter = Converter(str(source))

        try:
            converter.convert(str(target))
        finally:
            converter.close()

        if not target.is_file():
            raise RuntimeError(
                "Word conversion did not produce an output file."
            )

        return target.read_bytes()


# ---------------------------------------------------------------------------
# PDF to PowerPoint
# ---------------------------------------------------------------------------

def convert_pdf_to_pptx(pdf_bytes: bytes) -> bytes:
    """Create one image-based PowerPoint slide per PDF page."""
    with fitz.open(stream=pdf_bytes, filetype="pdf") as pdf:
        if pdf.page_count > MAX_PDF_PAGES:
            raise HTTPException(
                status_code=400,
                detail="PowerPoint conversion supports up to 200 pages.",
            )

        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        blank_layout = presentation.slide_layouts[6]

        for page in pdf:
            pixmap = page.get_pixmap(
                matrix=fitz.Matrix(1.5, 1.5),
                alpha=False,
            )

            image_stream = io.BytesIO(pixmap.tobytes("png"))
            slide = presentation.slides.add_slide(blank_layout)

            slide.shapes.add_picture(
                image_stream,
                0,
                0,
                width=presentation.slide_width,
                height=presentation.slide_height,
            )

        output = io.BytesIO()
        presentation.save(output)

        return output.getvalue()


# ---------------------------------------------------------------------------
# PDF to Excel
# ---------------------------------------------------------------------------

def convert_pdf_to_xlsx(pdf_bytes: bytes) -> bytes:
    """Extract PDF tables or text into an Excel workbook."""
    workbook = Workbook()
    default_sheet = workbook.active

    if default_sheet is None:
        raise RuntimeError("Could not create the Excel worksheet.")

    default_sheet.title = "Page 1"

    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        if len(pdf.pages) > 500:
            raise HTTPException(
                status_code=400,
                detail="Excel conversion supports up to 500 pages.",
            )

        for page_number, page in enumerate(pdf.pages, start=1):
            sheet_name = f"Page {page_number}"[:31]

            if page_number == 1:
                sheet = default_sheet
            else:
                sheet = workbook.create_sheet(title=sheet_name)

            tables = page.extract_tables() or []

            if tables:
                for table in tables:
                    for row in table:
                        sheet.append([
                            str(cell) if cell is not None else ""
                            for cell in row
                        ])

                    sheet.append([])
            else:
                page_text = page.extract_text() or ""

                for line in page_text.splitlines():
                    sheet.append([line])

    output = io.BytesIO()
    workbook.save(output)

    return output.getvalue()


# ---------------------------------------------------------------------------
# PDF to HTML
# ---------------------------------------------------------------------------

def convert_pdf_to_html(pdf_bytes: bytes) -> bytes:
    """
    Convert a PDF to standalone HTML.

    Preserves text positions, basic text styling, colors, page dimensions,
    and supported embedded raster images. Exact reproduction of every
    vector graphic, font metric, and PDF effect is not guaranteed.
    """
    page_sections: list[str] = []
    embedded_image_bytes = 0

    try:
        with fitz.open(
            stream=pdf_bytes,
            filetype="pdf",
        ) as document:

            if document.needs_pass:
                raise HTTPException(
                    status_code=400,
                    detail=(
                        "This PDF is password-protected. "
                        "Unlock it before converting."
                    ),
                )

            if document.page_count == 0:
                raise HTTPException(
                    status_code=400,
                    detail="The PDF contains no pages.",
                )

            if document.page_count > MAX_PDF_PAGES:
                raise HTTPException(
                    status_code=400,
                    detail="PDF to HTML supports up to 200 pages.",
                )

            for page_number, page in enumerate(document, start=1):
                page_width = max(float(page.rect.width), 1.0)
                page_height = max(float(page.rect.height), 1.0)

                positioned_elements: list[str] = []
                seen_image_rects: set[
                    tuple[int, int, int, int]
                ] = set()

                # Extract embedded raster images.
                for image_info in page.get_images(full=True):
                    xref = image_info[0]

                    if not xref:
                        continue

                    try:
                        image_data = document.extract_image(xref)
                        image_bytes = image_data.get("image", b"")
                        image_ext = image_data.get("ext", "").lower()

                        if not image_bytes:
                            continue

                        if (
                            embedded_image_bytes + len(image_bytes)
                            > MAX_HTML_IMAGE_BYTES
                        ):
                            logger.warning(
                                "HTML image limit reached on page %s",
                                page_number,
                            )
                            continue

                        allowed_image_types = {
                            "png": "image/png",
                            "jpg": "image/jpeg",
                            "jpeg": "image/jpeg",
                            "gif": "image/gif",
                            "webp": "image/webp",
                        }

                        mime_type = allowed_image_types.get(image_ext)

                        if not mime_type:
                            continue

                        image_data_uri = (
                            f"data:{mime_type};base64,"
                            + base64.b64encode(
                                image_bytes
                            ).decode("ascii")
                        )

                        rectangles = page.get_image_rects(xref)

                        for rect in rectangles:
                            rect_key = (
                                round(rect.x0),
                                round(rect.y0),
                                round(rect.x1),
                                round(rect.y1),
                            )

                            if rect_key in seen_image_rects:
                                continue

                            seen_image_rects.add(rect_key)

                            width = max(float(rect.width), 0.0)
                            height = max(float(rect.height), 0.0)

                            if width <= 0 or height <= 0:
                                continue

                            positioned_elements.append(
                                '<img class="pdf-image" '
                                f'src="{image_data_uri}" '
                                f'alt="Image from page {page_number}" '
                                'style="'
                                f'left:{rect.x0:.2f}px;'
                                f'top:{rect.y0:.2f}px;'
                                f'width:{width:.2f}px;'
                                f'height:{height:.2f}px;'
                                '">'
                            )

                        embedded_image_bytes += len(image_bytes)

                    except Exception as image_exc:
                        logger.debug(
                            "Could not extract image %s on page %s: %s",
                            xref,
                            page_number,
                            image_exc,
                        )

                # Extract text spans with their approximate formatting.
                page_dict = page.get_text("dict")

                for block in page_dict.get("blocks", []):
                    if block.get("type") != 0:
                        continue

                    for line in block.get("lines", []):
                        for span in line.get("spans", []):
                            text_value = str(span.get("text", ""))

                            if not text_value:
                                continue

                            bbox = span.get("bbox", (0, 0, 0, 0))

                            x0 = float(bbox[0])
                            y0 = float(bbox[1])
                            x1 = float(bbox[2])
                            y1 = float(bbox[3])

                            width = max(x1 - x0, 0.1)
                            height = max(y1 - y0, 0.1)

                            font_size = max(
                                1.0,
                                min(
                                    float(span.get("size", 11.0)),
                                    200.0,
                                ),
                            )

                            font_name = str(span.get("font", "Arial"))
                            font_family = css_font_family(font_name)
                            flags = int(span.get("flags", 0))

                            is_italic = bool(flags & 2)
                            is_bold = bool(flags & 16)
                            color = rgb_to_hex(
                                int(span.get("color", 0))
                            )

                            escaped_text = html_lib.escape(text_value)

                            styles = [
                                f"left:{x0:.2f}px",
                                f"top:{y0:.2f}px",
                                f"width:{width:.2f}px",
                                f"min-height:{height:.2f}px",
                                f"font-size:{font_size:.2f}px",
                                f"font-family:{font_family}",
                                f"color:{color}",
                                "white-space:pre",
                            ]

                            if is_bold:
                                styles.append("font-weight:700")

                            if is_italic:
                                styles.append("font-style:italic")

                            positioned_elements.append(
                                '<span class="pdf-text" '
                                f'style="{";".join(styles)}">'
                                f"{escaped_text}</span>"
                            )

                page_sections.append(
                    '<section class="pdf-page" '
                    f'aria-label="Page {page_number}" '
                    f'style="width:{page_width:.2f}px;'
                    f'height:{page_height:.2f}px;">'
                    f'<div class="page-label">Page {page_number}</div>'
                    + "".join(positioned_elements)
                    + "</section>"
                )

    except HTTPException:
        raise
    except Exception as exc:
        logger.exception("PDF-to-HTML conversion failed")

        raise HTTPException(
            status_code=400,
            detail=(
                "Could not read this PDF. "
                "It may be damaged or invalid."
            ),
        ) from exc

    html_document = """<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Converted PDF</title>
<style>
* {
  box-sizing: border-box;
}
html, body {
  margin: 0;
  padding: 0;
}
body {
  padding: 24px;
  background: #e5e7eb;
  color: #111827;
  font-family: Arial, Helvetica, sans-serif;
}
.pdf-page {
  position: relative;
  overflow: hidden;
  margin: 0 auto 24px;
  background: #fff;
  box-shadow: 0 2px 12px rgba(0, 0, 0, .12);
  page-break-after: always;
  break-after: page;
}
.pdf-text {
  position: absolute;
  display: block;
  z-index: 2;
  line-height: 1.15;
  transform-origin: top left;
}
.pdf-image {
  position: absolute;
  z-index: 1;
  display: block;
  object-fit: fill;
}
.page-label {
  position: absolute;
  right: 8px;
  bottom: 5px;
  z-index: 3;
  color: #9ca3af;
  font: 9px Arial, sans-serif;
}
@media print {
  @page {
    margin: 0;
  }
  body {
    padding: 0;
    background: #fff;
  }
  .pdf-page {
    margin: 0;
    box-shadow: none;
  }
}
</style>
</head>
<body>
""" + "\n".join(page_sections) + """
</body>
</html>
"""

    return html_document.encode("utf-8")


# ---------------------------------------------------------------------------
# HTML to PDF using Chrome or Edge
# ---------------------------------------------------------------------------

def convert_html_to_pdf_browser(
    file_bytes: bytes,
    filename: str,
) -> bytes:
    """
    Render an HTML document using headless Chrome or Edge.

    This generally handles modern HTML/CSS better than LibreOffice.
    The browser prints according to the document's print CSS and its
    default page configuration.

    JavaScript-generated content is not guaranteed to finish rendering
    before command-line printing. Relative assets must exist at their
    expected locations to load successfully.
    """
    browser = find_html_pdf_browser()

    if not browser:
        raise HTTPException(
            status_code=503,
            detail=(
                "Google Chrome or Microsoft Edge was not found. "
                "Install a supported browser, set HTML_PDF_BROWSER_PATH, "
                "or install LibreOffice for fallback conversion."
            ),
        )

    with tempfile.TemporaryDirectory() as directory:
        work_dir = Path(directory)
        source = work_dir / "input.html"
        output = work_dir / "output.pdf"
        profile = work_dir / "browser-profile"

        source.write_bytes(file_bytes)

        command = [
            browser,
            "--headless",
            "--disable-gpu",
            "--no-first-run",
            "--no-default-browser-check",
            "--disable-extensions",
            "--no-pdf-header-footer",
            f"--user-data-dir={profile}",
            f"--print-to-pdf={output}",
            source.resolve().as_uri(),
        ]

        try:
            result = subprocess.run(
                command,
                capture_output=True,
                text=True,
                timeout=120,
                check=False,
            )
        except subprocess.TimeoutExpired as exc:
            raise HTTPException(
                status_code=504,
                detail="HTML-to-PDF browser rendering timed out.",
            ) from exc

        if result.returncode != 0 or not output.is_file():
            logger.error(
                "Browser HTML-to-PDF failed. stdout=%s stderr=%s",
                result.stdout[-2000:],
                result.stderr[-2000:],
            )

            raise HTTPException(
                status_code=500,
                detail=(
                    "The browser could not render this HTML file. "
                    "Check the HTML, browser installation, and backend logs."
                ),
            )

        pdf_bytes = output.read_bytes()

        if not pdf_bytes.startswith(b"%PDF-"):
            raise HTTPException(
                status_code=500,
                detail="The browser returned an invalid PDF file.",
            )

        return pdf_bytes


# ---------------------------------------------------------------------------
# Office documents and HTML to PDF
# ---------------------------------------------------------------------------

def convert_office_to_pdf(
    file_bytes: bytes,
    filename: str,
    tool_slug: str | None = None,
) -> bytes:
    """
    Convert Office documents using LibreOffice.

    HTML-to-PDF uses Chrome/Edge first when available. LibreOffice is
    retained as a fallback if no supported browser is installed.
    """
    extension = Path(filename).suffix.lower()

    allowed_extensions = {
        ".docx",
        ".pptx",
        ".xlsx",
        ".html",
        ".htm",
    }

    if extension not in allowed_extensions:
        raise HTTPException(
            status_code=400,
            detail="Upload a DOCX, PPTX, XLSX, HTML, or HTM file.",
        )

    # Prefer a modern browser specifically for HTML-to-PDF.
    if tool_slug == "html-to-pdf":
        browser = find_html_pdf_browser()

        if browser:
            return convert_html_to_pdf_browser(
                file_bytes,
                filename,
            )

        logger.warning(
            "Chrome/Edge was not found. Falling back to LibreOffice "
            "for HTML-to-PDF."
        )

    libreoffice = find_libreoffice()

    if not libreoffice:
        raise HTTPException(
            status_code=503,
            detail=(
                "LibreOffice is not installed or could not be found. "
                "Install it and configure LIBREOFFICE_PATH if necessary."
            ),
        )

    with tempfile.TemporaryDirectory() as directory:
        work_dir = Path(directory)
        source = work_dir / f"input{extension}"
        output_dir = work_dir / "output"
        output_dir.mkdir()

        source.write_bytes(file_bytes)

        command = [
            libreoffice,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output_dir),
            str(source),
        ]

        try:
            result = subprocess.run(
                command,
                capture_output=True,
                text=True,
                timeout=180,
                check=False,
            )
        except subprocess.TimeoutExpired as exc:
            raise HTTPException(
                status_code=504,
                detail="Document-to-PDF conversion timed out.",
            ) from exc

        output_file = output_dir / "input.pdf"

        if result.returncode != 0 or not output_file.is_file():
            logger.error(
                "LibreOffice conversion failed. stdout=%s stderr=%s",
                result.stdout[-2000:],
                result.stderr[-2000:],
            )

            raise HTTPException(
                status_code=500,
                detail="LibreOffice could not convert this document to PDF.",
            )

        pdf_bytes = output_file.read_bytes()

        if not pdf_bytes.startswith(b"%PDF-"):
            raise HTTPException(
                status_code=500,
                detail="The rendering engine returned an invalid PDF.",
            )

        return pdf_bytes


# ---------------------------------------------------------------------------
# Health and root endpoints
# ---------------------------------------------------------------------------

@app.get("/")
def root() -> dict[str, str]:
    return {
        "message": "PDF Toolkit API is running.",
        "health": "/health",
        "docs": "/docs",
        "translation_languages": "/translation-languages",
    }


@app.get("/health")
def health() -> dict[str, Any]:
    return {
        "status": "ok",
        "service": "PDF Toolkit API",
        "libreoffice_available": find_libreoffice() is not None,
        "html_pdf_browser_available": find_html_pdf_browser() is not None,
        "translation_provider_url": os.getenv(
            "LIBRETRANSLATE_URL",
            "http://127.0.0.1:5000",
        ).rstrip("/"),
        "translation_routes_registered": True,
    }


# ---------------------------------------------------------------------------
# Document conversion endpoint
# ---------------------------------------------------------------------------

@app.post("/convert/{tool_slug}")
async def convert_document(
    tool_slug: str,
    file: UploadFile = File(...),
) -> Response:
    """Handle supported PDF and Office document conversions."""
    slug = validate_conversion_slug(tool_slug)
    filename = file.filename or "document"

    try:
        file_bytes = await read_upload(file)
        extension = Path(filename).suffix.lower()

        # ---------------------------------------------------------------
        # PDF-to-* conversions
        # ---------------------------------------------------------------
        if slug.startswith("pdf-to-"):
            if extension != ".pdf":
                raise HTTPException(
                    status_code=400,
                    detail="Please upload a PDF file for this conversion.",
                )

            require_pdf(file_bytes)

            if slug == "pdf-to-word":
                output = convert_pdf_to_docx(file_bytes)

                return file_response(
                    output,
                    f"{safe_stem(filename)}.docx",
                    (
                        "application/vnd.openxmlformats-officedocument."
                        "wordprocessingml.document"
                    ),
                )

            if slug == "pdf-to-powerpoint":
                output = convert_pdf_to_pptx(file_bytes)

                return file_response(
                    output,
                    f"{safe_stem(filename)}.pptx",
                    (
                        "application/vnd.openxmlformats-officedocument."
                        "presentationml.presentation"
                    ),
                )

            if slug == "pdf-to-excel":
                output = convert_pdf_to_xlsx(file_bytes)

                return file_response(
                    output,
                    f"{safe_stem(filename)}.xlsx",
                    (
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    ),
                )

            if slug == "pdf-to-html":
                output = convert_pdf_to_html(file_bytes)

                return file_response(
                    output,
                    f"{safe_stem(filename)}.html",
                    "text/html; charset=utf-8",
                )

        # ---------------------------------------------------------------
        # Office/HTML-to-PDF conversions
        # ---------------------------------------------------------------
        expected_extensions = {
            "word-to-pdf": {".docx"},
            "powerpoint-to-pdf": {".pptx"},
            "excel-to-pdf": {".xlsx"},
            "html-to-pdf": {".html", ".htm"},
        }

        if slug in expected_extensions:
            if extension not in expected_extensions[slug]:
                expected = ", ".join(
                    sorted(expected_extensions[slug])
                )

                raise HTTPException(
                    status_code=400,
                    detail=(
                        f"This tool requires a file extension: {expected}"
                    ),
                )

            output = convert_office_to_pdf(
                file_bytes,
                filename,
                tool_slug=slug,
            )

            return file_response(
                output,
                f"{safe_stem(filename)}.pdf",
                "application/pdf",
            )

        raise HTTPException(
            status_code=404,
            detail="Conversion tool not found.",
        )

    except HTTPException:
        raise

    except Exception as exc:
        logger.exception(
            "Conversion failed for tool %s",
            slug,
        )

        raise HTTPException(
            status_code=500,
            detail=(
                "Conversion failed. Check the file and backend logs."
            ),
        ) from exc

    finally:
        await file.close()


# ---------------------------------------------------------------------------
# Translation endpoints
# ---------------------------------------------------------------------------

app.include_router(translation_router)
app.include_router(ai_summarizer_router)

# Start the API with:
# uvicorn main:app --reload --host 127.0.0.1 --port 8000
